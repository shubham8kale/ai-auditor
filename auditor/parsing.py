"""Read spreadsheet cells without executing formulas or trusting footer totals."""

import csv
import io
import re
import zipfile
from datetime import date, datetime
from decimal import Decimal

from openpyxl import load_workbook

from auditor.domain import amount, finding, money, normalized

HEADERS = {
    "number": {
        "acct",
        "acct #",
        "acct no",
        "account",
        "account #",
        "account number",
        "account code",
        "account_number",
        "gl account",
        "code",
    },
    "name": {"account description", "account name", "account_name", "account title", "name"},
    "debit": {"debit", "debits", "dr"},
    "credit": {"credit", "credits", "cr"},
    "balance": {"balance", "net balance", "ending balance", "amount", "net amount"},
    "date": {"date", "posting date", "transaction date", "invoice date"},
    "counterparty": {"counterparty", "vendor", "supplier", "payee", "vendor name"},
    "ref": {"ref", "reference", "invoice number", "invoice #", "document number", "transaction id"},
    "description": {"description", "memo", "details", "transaction description"},
    "fsli": {"fsli", "financial statement line item", "category", "mapping"},
}
HEADERS = {key: {normalized(x) for x in values} for key, values in HEADERS.items()}


def date_string(value):
    if isinstance(value, (date, datetime)):
        return value.strftime("%Y-%m-%d")
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y", "%Y/%m/%d", "%Y-%m-%d %H:%M:%S", "%B %d %Y", "%B %d, %Y"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            continue
    raise ValueError(f"Unrecognized date: {text[:60]}")


def workbook_tables(data, extension):
    if extension == ".csv":
        text = data.decode("utf-8-sig")
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        return [
            {
                "name": "CSV",
                "rows": list(csv.reader(io.StringIO(text), dialect)),
                "formulas": [],
                "merged": [],
            }
        ]
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        if sum(item.file_size for item in archive.infolist()) > 80_000_000:
            raise ValueError("Expanded workbook exceeds the processing limit")
    workbook = load_workbook(io.BytesIO(data), read_only=False, data_only=False, keep_links=False)
    cached = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
    result = []
    try:
        for sheet in workbook:
            if sheet.max_row > 20000 or sheet.max_column > 100:
                raise ValueError("Workbook exceeds the supported 20,000-row / 100-column limit per sheet")
            formulas = [cell.coordinate for row in sheet for cell in row if cell.data_type == "f"]
            result.append(
                {
                    "name": sheet.title,
                    "rows": list(cached[sheet.title].values),
                    "formulas": formulas,
                    "merged": [str(r) for r in sheet.merged_cells.ranges],
                }
            )
    finally:
        workbook.close()
        cached.close()
    return result


def header_map(row):
    result = {}
    for index, value in enumerate(row):
        key = normalized(value or "")
        for target, aliases in HEADERS.items():
            if key in aliases and target not in result:
                result[target] = index
    return result


def spreadsheet(data: bytes, extension: str, document_id: str):
    tables = workbook_tables(data, extension)
    result = {
        "accounts": [],
        "transactions": [],
        "mappings": [],
        "sheets": [],
        "findings": [],
        "metadata_text": "",
    }
    detected_types = set()
    for table in tables:
        rows = table["rows"]
        result["sheets"].append(
            {
                "name": table["name"],
                "rows": len(rows),
                "formulas": table["formulas"],
                "merged": table["merged"],
            }
        )
        header = None
        kind = None
        for index, row in enumerate(rows[:40]):
            columns = header_map(row)
            if {"name", "fsli"} <= columns.keys():
                header, kind = (index, columns), "mapping_library"
                break
            if (
                "number" in columns
                and "date" in columns
                and ("balance" in columns or {"debit", "credit"} <= columns.keys())
            ):
                header, kind = (index, columns), "general_ledger"
                break
            if {"number", "name"} <= columns.keys() and (
                "balance" in columns or {"debit", "credit"} <= columns.keys()
            ):
                header, kind = (index, columns), "trial_balance"
                break
        if header is None:
            result["metadata_text"] += "\n".join(
                " | ".join(str(c) for c in row if c is not None) for row in rows[:100]
            )[:18000]
            continue
        detected_types.add(kind)
        index, columns = header
        result["metadata_text"] += (
            "\n".join(" | ".join(str(c) for c in row if c is not None) for row in rows[:index]) + "\n"
        )

        def get(row, key):
            return row[columns[key]] if key in columns and columns[key] < len(row) else None

        for row_index, row in enumerate(rows[index + 1 :], index + 2):
            if not any(value is not None and str(value).strip() for value in row):
                continue
            locator = f"{table['name']}!A{row_index}"
            number = get(row, "number")
            name = str(get(row, "name") or "").strip()
            if any(
                normalized(value or "").startswith(("total", "subtotal", "grand total")) for value in row[:2]
            ):
                continue
            if kind == "mapping_library":
                if name and get(row, "fsli"):
                    result["mappings"].append(
                        {"account_name": name, "fsli": str(get(row, "fsli")), "source": locator}
                    )
                continue
            if number is None or str(number).strip() == "":
                result["findings"].append(
                    finding(
                        "missing_account_number", "A populated row has no account identifier.", source=locator
                    )
                )
                continue
            number = (
                str(int(number)) if isinstance(number, float) and number.is_integer() else str(number).strip()
            )
            try:
                if "balance" in columns:
                    balance = money(get(row, "balance"))
                else:
                    if get(row, "debit") in (None, "") and get(row, "credit") in (None, ""):
                        raise ValueError("Both debit and credit are missing; an empty amount is not zero")
                    debit = money(get(row, "debit") or 0)
                    credit = money(get(row, "credit") or 0)
                    balance = debit - credit
                if kind == "trial_balance":
                    if not name:
                        raise ValueError("Account name is missing")
                    result["accounts"].append(
                        {
                            "number": number,
                            "name": name,
                            "balance": amount(balance),
                            "source": locator,
                            "document_id": document_id,
                        }
                    )
                else:
                    parsed_date = date_string(get(row, "date"))
                    result["transactions"].append(
                        {
                            "id": f"{document_id}:{table['name']}:{row_index}",
                            "account_number": number,
                            "account_name": name,
                            "date": parsed_date,
                            "counterparty": str(get(row, "counterparty") or ""),
                            "ref": str(get(row, "ref") or ""),
                            "description": str(get(row, "description") or ""),
                            "amount": amount(balance),
                            "source": locator,
                            "document_id": document_id,
                        }
                    )
            except ValueError as exc:
                result["findings"].append(finding("invalid_row", str(exc), source=locator))
    result["kind"] = next(iter(detected_types)) if len(detected_types) == 1 else "other_support"
    if len(detected_types) > 1:
        result["findings"].append(
            finding(
                "mixed_types",
                "This workbook contains multiple document types. Split the types into separate uploads.",
            )
        )
    if result["accounts"]:
        total = sum((money(row["balance"]) for row in result["accounts"]), Decimal(0))
        result["net_balance"] = amount(total)
        if total:
            result["findings"].append(
                finding(
                    "tb_unbalanced",
                    f"Trial balance is out of balance by {amount(total)}.",
                    source="POL-103 §3",
                )
            )
        identifiers = [row["number"] for row in result["accounts"]]
        if len(set(identifiers)) != len(identifiers):
            result["findings"].append(
                finding(
                    "duplicate_account",
                    "Account identifiers repeat. Supply one row per account or explain the source grouping.",
                    source="POL-103 §3",
                )
            )
    result["has_formulas"] = any(s["formulas"] for s in result["sheets"])
    result["has_export_header"] = bool(
        re.search(r"export|quickbooks|xero|netsuite|sage|generated by", result["metadata_text"], re.I)
    )
    result["source_rows"] = [
        {
            "sheet": t["name"],
            "rows": [
                [str(v) if isinstance(v, (date, datetime, Decimal)) else v for v in row]
                for row in t["rows"][:300]
            ],
        }
        for t in tables
    ]
    return result
